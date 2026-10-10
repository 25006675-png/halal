"""Excel and Word files turned into text for Gemini. Run: python -m pytest tests"""

import time
from datetime import datetime
from io import BytesIO

import openpyxl
from docx import Document

from agents import office


def xlsx_bytes(rows: list[list], title: str = "Ingredients") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Receiving procedure", level=1)
    doc.add_paragraph("Check every delivery against the approved supplier list.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Supplier", "Brand"
    table.cell(1, 0).text, table.cell(1, 1).text = "Minyak Murni Sdn Bhd", "MURNI"
    doc.add_heading("Records", level=2)
    doc.add_paragraph("Keep the receiving log for two years.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_kind_by_extension_or_type():
    assert office.kind(None, "list.xlsx") == "xlsx"
    assert office.kind("application/octet-stream", "LIST.XLSM") == "xlsx"
    assert office.kind("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "upload") == "docx"
    assert office.kind(None, "old.xls") == "legacy"
    assert office.kind(None, "old.doc") == "legacy"
    assert office.kind("image/png", "photo.png") is None
    assert office.kind("application/pdf", "cert.pdf") is None


def test_xlsx_rows_dates_and_blanks():
    data = xlsx_bytes([
        ["Ingredient", "Brand", "Supplier", "Expiry"],
        ["Cooking oil", "MURNI", "Minyak Murni Sdn Bhd", datetime(2027, 6, 30)],
        [None, None, None, None],                       # blank row: skipped
        ["Sugar", "GULA MANIS", "Gula Manis Sdn Bhd", 12.0],
    ])
    text = office.to_text(data, None, "ingredient_list.xlsx")
    assert text.splitlines() == [
        "## Sheet: Ingredients",
        "Ingredient | Brand | Supplier | Expiry",
        "Cooking oil | MURNI | Minyak Murni Sdn Bhd | 2027-06-30",
        "Sugar | GULA MANIS | Gula Manis Sdn Bhd | 12",
    ]


def test_docx_keeps_order_headings_and_tables():
    text = office.to_text(docx_bytes(), None, "sop.docx")
    assert text.splitlines() == [
        "# Receiving procedure",
        "Check every delivery against the approved supplier list.",
        "Supplier | Brand",
        "Minyak Murni Sdn Bhd | MURNI",
        "## Records",
        "Keep the receiving log for two years.",
    ]


def test_not_converted_for_other_files():
    assert office.to_text(b"\x89PNG", "image/png", "photo.png") is None


def test_fast_enough():
    data = xlsx_bytes([["Item", "Brand", "Qty"]] + [[f"Item {i}", "BRAND", i] for i in range(500)])
    t = time.perf_counter()
    office.to_text(data, None, "big.xlsx")
    office.to_text(docx_bytes(), None, "sop.docx")
    assert time.perf_counter() - t < 2.0
