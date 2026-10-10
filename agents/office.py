"""
Excel and Word files as text, for Gemini (which takes images, PDFs and text,
but not .xlsx or .docx). Pure Python and fast: well under a second for files
the size a small business has.
"""

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

MAX_ROWS = 2000          # per sheet; an ingredient list is far smaller
MAX_CHARS = 60000        # what we send to Gemini

XLSX_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              "application/vnd.ms-excel.sheet.macroenabled.12"}
DOCX_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def kind(content_type: str | None, filename: str) -> str | None:
    """'xlsx', 'docx', 'legacy' (old .xls/.doc), or None for anything else."""
    ext = Path(filename or "").suffix.lower()
    ctype = (content_type or "").lower()
    if ext in (".xlsx", ".xlsm") or ctype in XLSX_TYPES:
        return "xlsx"
    if ext == ".docx" or ctype in DOCX_TYPES:
        return "docx"
    if ext in (".xls", ".doc") or ctype in ("application/vnd.ms-excel", "application/msword"):
        return "legacy"
    return None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def xlsx_text(data: bytes) -> str:
    import openpyxl
    # data_only: the values Excel last saved, not the formulas
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            rows = 0
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if rows == MAX_ROWS:
                    out.append(f"(more rows not shown; first {MAX_ROWS} only)")
                    break
                out.append(" | ".join(cells))
                rows += 1
    finally:
        wb.close()
    return "\n".join(out)[:MAX_CHARS]


def docx_text(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(BytesIO(data))
    out = []
    # Paragraphs and tables in the order they appear in the document.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading") and style[-1:].isdigit():
                out.append("#" * min(4, int(style[-1])) + " " + text)
            elif style.lower() == "title":
                out.append("# " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    images = len(doc.inline_shapes)
    if images:
        out.append(f"(The document also contains {images} picture(s) that were not read.)")
    return "\n".join(out)[:MAX_CHARS]


def to_text(data: bytes, content_type: str | None, filename: str) -> str | None:
    """Text for an Excel or Word file; None if it isn't one we convert."""
    k = kind(content_type, filename)
    if k == "xlsx":
        return xlsx_text(data)
    if k == "docx":
        return docx_text(data)
    return None
